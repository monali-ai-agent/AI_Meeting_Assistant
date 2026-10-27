"""Tools: Document parser, Task tool, due-date parsing, team directory."""
import io
import re
from datetime import date, timedelta

DEFAULT_TEAM = [  # add each person's real email address
    {"name": "Priya", "email": "", "role": "Product Manager", "skills": ["roadmap", "launch", "product", "release", "beta", "requirements"]},
    {"name": "Rahul", "email": "", "role": "Business Analyst", "skills": ["pricing", "deck", "budget", "report", "finance", "presentation"]},
    {"name": "Anita", "email": "", "role": "QA Lead", "skills": ["qa", "test", "testing", "bug", "quality", "regression"]},
    {"name": "Karthik", "email": "", "role": "Backend Engineer", "skills": ["api", "backend", "database", "server", "integration", "deploy"]},
    {"name": "Sneha", "email": "", "role": "Designer", "skills": ["design", "ui", "ux", "mockup", "branding", "figma"]},
]

WEEKDAYS = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]


def parse_due(text, today=None):
    """Turn phrases like 'by Friday', 'Nov 15', 'tomorrow' into an ISO date (or None)."""
    today = today or date.today()
    t = (text or "").lower()
    m = re.search(r"\b(\d{4}-\d{2}-\d{2})\b", t)
    if m:
        return m.group(1)
    if "tomorrow" in t:
        return (today + timedelta(days=1)).isoformat()
    if re.search(r"\b(today|eod|end of day)\b", t):
        return today.isoformat()
    m = re.search(r"\b(next\s+)?(monday|tuesday|wednesday|thursday|friday|saturday|sunday)\b", t)
    if m:
        ahead = (WEEKDAYS.index(m.group(2)) - today.weekday()) % 7 or 7
        return (today + timedelta(days=ahead)).isoformat()
    m = re.search(r"\b(" + "|".join(MONTHS) + r")[a-z]*\.?\s+(\d{1,2})(?:st|nd|rd|th)?\b", t) or None
    if m:
        mon, day = MONTHS.index(m.group(1)) + 1, int(m.group(2))
    else:
        m2 = re.search(r"\b(\d{1,2})(?:st|nd|rd|th)?\s+(?:of\s+)?(" + "|".join(MONTHS) + r")[a-z]*\b", t)
        if not m2:
            if re.search(r"\b(next week)\b", t):
                return (today + timedelta(days=7)).isoformat()
            if re.search(r"\b(end of (the )?week|this week)\b", t):
                return (today + timedelta(days=(4 - today.weekday()) % 7)).isoformat()
            return None
        day, mon = int(m2.group(1)), MONTHS.index(m2.group(2)) + 1
    try:
        d = date(today.year, mon, day)
        if d < today:
            d = date(today.year + 1, mon, day)
        return d.isoformat()
    except ValueError:
        return None


class DocumentParser:
    """Reads TXT/MD, DOCX, PDF, VTT/SRT into plain text with 'Speaker: text' lines when possible."""

    def parse(self, filename, data: bytes) -> str:
        ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else "txt"
        if ext == "docx":
            return self._docx(data)
        if ext == "pdf":
            return self._pdf(data)
        if ext in ("vtt", "srt"):
            return self._subtitles(data.decode("utf-8", errors="replace"))
        return data.decode("utf-8", errors="replace")

    def _docx(self, data):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for tbl in d.tables:
            for row in tbl.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)

    def _pdf(self, data):
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in r.pages)

    def _subtitles(self, text):
        out = []
        for line in text.splitlines():
            line = line.strip()
            if not line or line == "WEBVTT" or re.match(r"^\d+$", line) or "-->" in line:
                continue
            line = re.sub(r"<v\s+([^>]+)>", r"\1: ", line)
            line = re.sub(r"</?[^>]+>", "", line)
            out.append(line)
        return "\n".join(out)


class TaskTool:
    """Task tool used by the Assignment and Summary agents (create / update / list / remind)."""

    def __init__(self, memory):
        self.mem = memory

    def create_task(self, meeting_id, task, owner, due, priority="medium"):
        return self.mem.add_task(meeting_id, task, owner, due, priority)

    def update_status(self, task_id, status):
        return self.mem.update_task(task_id, status=status)

    def assign(self, task_id, owner):
        return self.mem.update_task(task_id, owner=owner)

    def list_pending(self, owner=None):
        return self.mem.list_tasks(pending_only=True, owner=owner)

    def overdue(self):
        today = date.today().isoformat()
        return [t for t in self.list_pending() if t["due"] and t["due"] < today]

    def send_reminder(self, task):
        """Simulated reminder. Replace with email / Slack / Teams in production."""
        who = task["owner"] or "the team"
        due = f" (due {task['due']})" if task["due"] else ""
        return f"Reminder for {who}: task #{task['id']} \"{task['task']}\"{due}"
